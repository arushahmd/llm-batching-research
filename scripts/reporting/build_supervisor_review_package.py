from __future__ import annotations

import html
import re
import shutil
from dataclasses import dataclass
from pathlib import Path

import fitz
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt
from PIL import Image
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, StyleSheet1
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Image as RLImage
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


PROJECT_ROOT = Path(__file__).resolve().parents[2]
SOURCE_MD = PROJECT_ROOT / "final_corrected_thesis" / "master_chapters_1_to_5_corrected.md"
FIGURE_INDEX_MD = PROJECT_ROOT / "reports" / "thesis_figures" / "FIGURE_INDEX.md"
SOURCE_FIGURES_DIR = PROJECT_ROOT / "reports" / "thesis_figures"
OUTPUT_DIR = PROJECT_ROOT / "Supervisor_Review_Package"
OUTPUT_FIGURES_DIR = OUTPUT_DIR / "figures"
DOCX_PATH = OUTPUT_DIR / "thesis_chapters_1_to_5_review.docx"
PDF_PATH = OUTPUT_DIR / "thesis_chapters_1_to_5_review.pdf"
README_PATH = OUTPUT_DIR / "README_REVIEW_DRAFT.md"


@dataclass
class FigureInfo:
    number: str
    source_path: Path
    package_path: Path
    caption: str


@dataclass
class Block:
    kind: str
    data: dict


def read_text(path: Path) -> str:
    return path.read_text(encoding="utf-8")


def clean_inline_text(text: str) -> str:
    return text.replace("`", "").strip()


def is_block_start(line: str) -> bool:
    stripped = line.strip()
    if not stripped:
        return True
    patterns = [
        r"^Chapter\s+\d+$",
        r"^##\s+",
        r"^###\s+",
        r"^\[Insert Figure\s+\d+\.\d+\s+here:",
        r"^Figure\s+\d+\.\d+\.",
        r"^Table\s+\d+\.\d+\.",
        r"^- ",
        r"^\d+\.\s+",
        r"^\|",
    ]
    return any(re.match(pattern, stripped) for pattern in patterns)


def parse_figure_index(path: Path) -> dict[str, FigureInfo]:
    figure_map: dict[str, FigureInfo] = {}
    for line in read_text(path).splitlines():
        if not line.startswith("| Figure "):
            continue
        columns = [part.strip() for part in line.strip().strip("|").split("|")]
        if len(columns) < 7:
            continue
        number = columns[0]
        file_match = re.search(r"`([^`]+)`", columns[1])
        if not file_match:
            continue
        relative_path = Path(file_match.group(1))
        source_path = PROJECT_ROOT / relative_path
        package_path = OUTPUT_FIGURES_DIR / source_path.name
        caption = columns[4]
        figure_map[number] = FigureInfo(
            number=number,
            source_path=source_path,
            package_path=package_path,
            caption=caption,
        )
    return figure_map


def parse_markdown_table(lines: list[str]) -> list[list[str]]:
    rows: list[list[str]] = []
    for index, line in enumerate(lines):
        if index == 1:
            continue
        raw_cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
        rows.append([clean_inline_text(cell) for cell in raw_cells])
    return rows


def parse_markdown_blocks(text: str, figure_map: dict[str, FigureInfo]) -> list[Block]:
    lines = text.splitlines()
    blocks: list[Block] = []
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if not stripped:
            i += 1
            continue

        if re.match(r"^Chapter\s+\d+$", stripped):
            chapter_label = stripped
            i += 1
            while i < len(lines) and not lines[i].strip():
                i += 1
            title = lines[i].strip() if i < len(lines) else ""
            chapter_number = chapter_label.split()[1]
            blocks.append(
                Block(
                    kind="chapter",
                    data={
                        "number": chapter_number,
                        "title": title,
                        "display": f"{chapter_label}: {title}",
                    },
                )
            )
            i += 1
            continue

        if stripped.startswith("## "):
            blocks.append(Block(kind="heading2", data={"text": clean_inline_text(stripped[3:])}))
            i += 1
            continue

        if stripped.startswith("### "):
            blocks.append(Block(kind="heading3", data={"text": clean_inline_text(stripped[4:])}))
            i += 1
            continue

        figure_match = re.match(r"^\[Insert Figure\s+(\d+\.\d+)\s+here:", stripped)
        if figure_match:
            number = f"Figure {figure_match.group(1)}"
            if number not in figure_map:
                raise KeyError(f"Figure not found in index: {number}")
            blocks.append(Block(kind="figure", data={"figure": figure_map[number]}))
            i += 1
            while i < len(lines) and not lines[i].strip():
                i += 1
            if i < len(lines) and lines[i].strip().startswith(f"{number}."):
                i += 1
            continue

        if re.match(r"^Figure\s+\d+\.\d+\.", stripped):
            i += 1
            continue

        if re.match(r"^Table\s+\d+\.\d+\.", stripped):
            caption = clean_inline_text(stripped)
            i += 1
            while i < len(lines) and not lines[i].strip():
                i += 1
            table_lines: list[str] = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i])
                i += 1
            blocks.append(
                Block(
                    kind="table",
                    data={
                        "caption": caption,
                        "rows": parse_markdown_table(table_lines),
                    },
                )
            )
            continue

        if stripped.startswith("- "):
            items: list[str] = []
            while i < len(lines):
                current = lines[i]
                current_stripped = current.strip()
                if not current_stripped:
                    i += 1
                    break
                if current_stripped.startswith("- "):
                    items.append(clean_inline_text(current_stripped[2:]))
                    i += 1
                    continue
                if current.startswith("  ") or current.startswith("\t"):
                    items[-1] = f"{items[-1]} {clean_inline_text(current_stripped)}"
                    i += 1
                    continue
                break
            blocks.append(Block(kind="bullet_list", data={"items": items}))
            continue

        if re.match(r"^\d+\.\s+", stripped):
            items: list[str] = []
            while i < len(lines):
                current = lines[i]
                current_stripped = current.strip()
                if not current_stripped:
                    i += 1
                    break
                match = re.match(r"^(\d+)\.\s+(.*)$", current_stripped)
                if match:
                    items.append(clean_inline_text(match.group(2)))
                    i += 1
                    continue
                if current.startswith("   ") or current.startswith("\t"):
                    items[-1] = f"{items[-1]} {clean_inline_text(current_stripped)}"
                    i += 1
                    continue
                break
            blocks.append(Block(kind="numbered_list", data={"items": items}))
            continue

        paragraph_lines = [stripped]
        i += 1
        while i < len(lines):
            current = lines[i]
            current_stripped = current.strip()
            if not current_stripped:
                i += 1
                break
            if is_block_start(current):
                break
            paragraph_lines.append(current_stripped)
            i += 1
        paragraph = clean_inline_text(" ".join(paragraph_lines))
        blocks.append(Block(kind="paragraph", data={"text": paragraph}))

    return blocks


def copy_figures(figure_map: dict[str, FigureInfo]) -> None:
    OUTPUT_FIGURES_DIR.mkdir(parents=True, exist_ok=True)
    copied = set()
    for info in figure_map.values():
        if info.source_path.name in copied:
            continue
        shutil.copy2(info.source_path, info.package_path)
        copied.add(info.source_path.name)


def image_dimensions(image_path: Path, max_width_in: float, max_height_in: float) -> tuple[float, float]:
    with Image.open(image_path) as image:
        width_px, height_px = image.size
    aspect = width_px / height_px
    width_in = max_width_in
    height_in = width_in / aspect
    if height_in > max_height_in:
        height_in = max_height_in
        width_in = height_in * aspect
    return width_in, height_in


def set_cell_text(cell, text: str, *, bold: bool = False, font_size: int = 10) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if bold else WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(font_size)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def build_docx(blocks: list[Block]) -> None:
    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    normal_style = document.styles["Normal"]
    normal_style.font.name = "Times New Roman"
    normal_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal_style.font.size = Pt(12)

    first_chapter = True
    for block in blocks:
        kind = block.kind
        data = block.data
        if kind == "chapter":
            if not first_chapter:
                document.add_page_break()
            first_chapter = False
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(18)
            run = paragraph.add_run(data["display"])
            run.bold = True
            run.font.name = "Times New Roman"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
            run.font.size = Pt(16)
            continue

        if kind == "heading2":
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(12)
            paragraph.paragraph_format.space_after = Pt(6)
            run = paragraph.add_run(data["text"])
            run.bold = True
            run.font.name = "Times New Roman"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
            run.font.size = Pt(14)
            continue

        if kind == "heading3":
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(10)
            paragraph.paragraph_format.space_after = Pt(4)
            run = paragraph.add_run(data["text"])
            run.bold = True
            run.font.name = "Times New Roman"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
            run.font.size = Pt(12)
            continue

        if kind == "paragraph":
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
            paragraph.paragraph_format.space_after = Pt(6)
            run = paragraph.add_run(data["text"])
            run.font.name = "Times New Roman"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
            run.font.size = Pt(12)
            continue

        if kind == "bullet_list":
            for item in data["items"]:
                paragraph = document.add_paragraph(style="List Bullet")
                paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
                paragraph.paragraph_format.space_after = Pt(3)
                run = paragraph.add_run(item)
                run.font.name = "Times New Roman"
                run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
                run.font.size = Pt(12)
            continue

        if kind == "numbered_list":
            for item in data["items"]:
                paragraph = document.add_paragraph(style="List Number")
                paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
                paragraph.paragraph_format.space_after = Pt(3)
                run = paragraph.add_run(item)
                run.font.name = "Times New Roman"
                run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
                run.font.size = Pt(12)
            continue

        if kind == "figure":
            figure: FigureInfo = data["figure"]
            max_width_in = 6.0
            max_height_in = 8.0
            width_in, _ = image_dimensions(figure.package_path, max_width_in, max_height_in)
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run()
            run.add_picture(str(figure.package_path), width=Inches(width_in))

            caption = document.add_paragraph()
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            caption.paragraph_format.space_after = Pt(8)
            cap_run = caption.add_run(f"{figure.number}. {figure.caption}")
            cap_run.italic = True
            cap_run.font.name = "Times New Roman"
            cap_run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
            cap_run.font.size = Pt(11)
            continue

        if kind == "table":
            caption = document.add_paragraph()
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            caption.paragraph_format.space_after = Pt(4)
            cap_run = caption.add_run(data["caption"])
            cap_run.bold = True
            cap_run.font.name = "Times New Roman"
            cap_run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
            cap_run.font.size = Pt(11)

            rows = data["rows"]
            table = document.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = True
            set_repeat_table_header(table.rows[0])
            for row_index, row in enumerate(rows):
                for col_index, cell_value in enumerate(row):
                    set_cell_text(
                        table.cell(row_index, col_index),
                        cell_value,
                        bold=row_index == 0,
                        font_size=9 if len(rows[0]) >= 8 else 10,
                    )
            document.add_paragraph()
            continue

    document.save(DOCX_PATH)


def register_pdf_fonts() -> tuple[str, str, str]:
    times_regular = Path(r"C:\Windows\Fonts\times.ttf")
    times_bold = Path(r"C:\Windows\Fonts\timesbd.ttf")
    times_italic = Path(r"C:\Windows\Fonts\timesi.ttf")
    if times_regular.exists() and times_bold.exists() and times_italic.exists():
        pdfmetrics.registerFont(TTFont("TimesNewRoman", str(times_regular)))
        pdfmetrics.registerFont(TTFont("TimesNewRoman-Bold", str(times_bold)))
        pdfmetrics.registerFont(TTFont("TimesNewRoman-Italic", str(times_italic)))
        return "TimesNewRoman", "TimesNewRoman-Bold", "TimesNewRoman-Italic"
    return "Times-Roman", "Times-Bold", "Times-Italic"


def build_pdf_styles(font_regular: str, font_bold: str, font_italic: str) -> StyleSheet1:
    styles = StyleSheet1()
    styles.add(
        ParagraphStyle(
            name="Body",
            fontName=font_regular,
            fontSize=12,
            leading=18,
            alignment=TA_JUSTIFY,
            spaceAfter=8,
        )
    )
    styles.add(
        ParagraphStyle(
            name="Chapter",
            fontName=font_bold,
            fontSize=16,
            leading=20,
            alignment=TA_CENTER,
            spaceAfter=16,
        )
    )
    styles.add(
        ParagraphStyle(
            name="Heading2",
            fontName=font_bold,
            fontSize=14,
            leading=18,
            alignment=TA_LEFT,
            spaceBefore=10,
            spaceAfter=6,
        )
    )
    styles.add(
        ParagraphStyle(
            name="Heading3",
            fontName=font_bold,
            fontSize=12,
            leading=16,
            alignment=TA_LEFT,
            spaceBefore=8,
            spaceAfter=4,
        )
    )
    styles.add(
        ParagraphStyle(
            name="Caption",
            fontName=font_italic,
            fontSize=11,
            leading=14,
            alignment=TA_CENTER,
            spaceAfter=10,
        )
    )
    styles.add(
        ParagraphStyle(
            name="TableCaption",
            fontName=font_bold,
            fontSize=11,
            leading=14,
            alignment=TA_CENTER,
            spaceAfter=5,
        )
    )
    styles.add(
        ParagraphStyle(
            name="List",
            fontName=font_regular,
            fontSize=12,
            leading=18,
            alignment=TA_JUSTIFY,
            leftIndent=18,
            firstLineIndent=-12,
            spaceAfter=4,
        )
    )
    styles.add(
        ParagraphStyle(
            name="TableCell",
            fontName=font_regular,
            fontSize=8.8,
            leading=10.5,
            alignment=TA_LEFT,
        )
    )
    styles.add(
        ParagraphStyle(
            name="TableHeader",
            parent=styles["TableCell"],
            fontName=font_bold,
            alignment=TA_CENTER,
        )
    )
    return styles


def pdf_paragraph(text: str, style: ParagraphStyle) -> Paragraph:
    return Paragraph(html.escape(text), style)


def build_pdf(blocks: list[Block]) -> None:
    font_regular, font_bold, font_italic = register_pdf_fonts()
    styles = build_pdf_styles(font_regular, font_bold, font_italic)
    doc = SimpleDocTemplate(
        str(PDF_PATH),
        pagesize=A4,
        leftMargin=inch,
        rightMargin=inch,
        topMargin=inch,
        bottomMargin=inch,
    )
    story = []
    first_chapter = True

    for block in blocks:
        kind = block.kind
        data = block.data

        if kind == "chapter":
            if not first_chapter:
                story.append(PageBreak())
            first_chapter = False
            story.append(pdf_paragraph(data["display"], styles["Chapter"]))
            continue

        if kind == "heading2":
            story.append(pdf_paragraph(data["text"], styles["Heading2"]))
            continue

        if kind == "heading3":
            story.append(pdf_paragraph(data["text"], styles["Heading3"]))
            continue

        if kind == "paragraph":
            story.append(pdf_paragraph(data["text"], styles["Body"]))
            continue

        if kind == "bullet_list":
            for item in data["items"]:
                story.append(pdf_paragraph(f"• {item}", styles["List"]))
            story.append(Spacer(1, 4))
            continue

        if kind == "numbered_list":
            for index, item in enumerate(data["items"], start=1):
                story.append(pdf_paragraph(f"{index}. {item}", styles["List"]))
            story.append(Spacer(1, 4))
            continue

        if kind == "figure":
            figure: FigureInfo = data["figure"]
            max_width_in, max_height_in = 6.0, 8.0
            width_in, height_in = image_dimensions(figure.package_path, max_width_in, max_height_in)
            image = RLImage(str(figure.package_path), width=width_in * inch, height=height_in * inch)
            image.hAlign = "CENTER"
            story.append(image)
            story.append(pdf_paragraph(f"{figure.number}. {figure.caption}", styles["Caption"]))
            continue

        if kind == "table":
            story.append(pdf_paragraph(data["caption"], styles["TableCaption"]))
            rows = data["rows"]
            wrapped_rows = []
            for row_index, row in enumerate(rows):
                wrapped_rows.append(
                    [
                        pdf_paragraph(cell, styles["TableHeader"] if row_index == 0 else styles["TableCell"])
                        for cell in row
                    ]
                )
            col_count = len(rows[0])
            usable_width = A4[0] - 2 * inch
            col_width = usable_width / col_count
            table = Table(wrapped_rows, colWidths=[col_width] * col_count, repeatRows=1)
            table.setStyle(
                TableStyle(
                    [
                        ("GRID", (0, 0), (-1, -1), 0.5, colors.black),
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EAEAEA")),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("LEFTPADDING", (0, 0), (-1, -1), 4),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                        ("TOPPADDING", (0, 0), (-1, -1), 3),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                    ]
                )
            )
            story.append(table)
            story.append(Spacer(1, 10))
            continue

    doc.build(story)


def write_readme() -> None:
    lines = [
        "# Supervisor Review Draft",
        "",
        "This is a supervisor-review draft.",
        "It includes Chapters 1 to 5 only.",
        "Front matter such as declaration, certificate, acknowledgements, TOC, list of figures, and list of tables will be added after supervisor approval.",
        "Figures are inserted from `reports/thesis_figures/`.",
        "Results and claims are based on the corrected thesis draft.",
    ]
    README_PATH.write_text("\n".join(lines) + "\n", encoding="utf-8")


def verify_docx() -> None:
    document = Document(DOCX_PATH)
    paragraph_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    if "[Insert Figure" in paragraph_text:
        raise ValueError("DOCX still contains figure placeholders.")
    if document.inline_shapes and len(document.inline_shapes) < 17:
        raise ValueError(f"DOCX contains too few figures: {len(document.inline_shapes)}")
    for required in [
        "Figure 1.1.",
        "Figure 5.5.",
        "Figure 5.6.",
        "Figure 5.7.",
    ]:
        if required not in paragraph_text:
            raise ValueError(f"DOCX is missing caption text: {required}")
    if "The first two points in Figure 5.5 are directly comparable under the same training budget." not in paragraph_text:
        raise ValueError("Figure 5.5 caveat text is missing from DOCX body text.")
    chapter4_order = [line for line in paragraph_text.splitlines() if line.startswith("Figure 4.")]
    expected_prefixes = ["Figure 4.1.", "Figure 4.2.", "Figure 4.3.", "Figure 4.4."]
    if chapter4_order[:4] != [next(text for text in chapter4_order if text.startswith(prefix)) for prefix in expected_prefixes]:
        raise ValueError("Chapter 4 figure order is not sequential in DOCX.")


def verify_pdf() -> None:
    pdf = fitz.open(PDF_PATH)
    if pdf.page_count <= 0:
        raise ValueError("PDF did not render any pages.")
    image_count = 0
    full_text_parts = []
    for page in pdf:
        image_count += len(page.get_images(full=True))
        full_text_parts.append(page.get_text("text"))
    pdf.close()
    full_text = "\n".join(full_text_parts)
    if "[Insert Figure" in full_text:
        raise ValueError("PDF still contains figure placeholders.")
    if image_count < 17:
        raise ValueError(f"PDF appears to contain too few images: {image_count}")
    for required in [
        "Figure 1.1.",
        "Figure 5.5.",
        "Figure 5.6.",
        "Figure 5.7.",
    ]:
        if required not in full_text:
            raise ValueError(f"PDF is missing caption text: {required}")
    if "The first two points in Figure 5.5 are directly comparable under the same training budget." not in full_text:
        raise ValueError("Figure 5.5 caveat text is missing from PDF body text.")
    chapter4_positions = [full_text.find(f"Figure 4.{idx}.") for idx in range(1, 5)]
    if any(position == -1 for position in chapter4_positions) or chapter4_positions != sorted(chapter4_positions):
        raise ValueError("Chapter 4 figures are not in sequential order in PDF.")


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    figure_map = parse_figure_index(FIGURE_INDEX_MD)
    copy_figures(figure_map)
    blocks = parse_markdown_blocks(read_text(SOURCE_MD), figure_map)
    build_docx(blocks)
    build_pdf(blocks)
    write_readme()
    verify_docx()
    verify_pdf()
    print(f"Created {DOCX_PATH}")
    print(f"Created {PDF_PATH}")
    print(f"Created {OUTPUT_FIGURES_DIR}")
    print(f"Created {README_PATH}")


if __name__ == "__main__":
    main()
